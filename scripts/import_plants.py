from __future__ import annotations

import json
import os
import re
from collections import Counter
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd
from PIL import Image, ImageOps


ROOT = Path(__file__).resolve().parents[1]
SOURCE_ROOT = Path(os.getenv("BASHANG_SOURCE_ROOT", ROOT / "source-data"))
WORKBOOK_PATH = SOURCE_ROOT / "河北丰宁坝上植物知识库_群聊图片已导入_20260710.xlsx"
PUBLIC_DATA_DIR = ROOT / "public" / "data"
REFERENCE_IMAGE_DIR = ROOT / "public" / "plant-images"
LOCAL_IMAGE_DIR = ROOT / "public" / "local-samples"

PLACEHOLDERS = {"", "20", "nan", "none", "null", "-", "—"}
IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".webp"}

FLOWER_FIELDS = [
    "花颜色",
    "花序类型",
    "花冠/花形",
    "花对称性",
    "花性",
    "花瓣/花被片数",
    "花大小",
    "花期",
    "花完整描述",
]
FRUIT_FIELDS = [
    "果实类型",
    "果实形状",
    "果实颜色",
    "果实大小",
    "果期",
    "种子特征",
    "果实完整描述",
]
STEM_FIELDS = [
    "茎质地",
    "茎姿态",
    "分枝方式",
    "茎形/结构",
    "茎颜色",
    "茎毛被/刺",
    "植株高度",
    "茎完整描述",
]
LEAF_FIELDS = [
    "叶序",
    "叶序细分",
    "叶类型",
    "复叶类型",
    "叶片形状",
    "叶缘",
    "叶尖",
    "叶基",
    "叶脉",
    "叶质",
    "叶面/毛被",
    "叶柄",
    "托叶/叶鞘",
    "叶大小",
    "叶完整描述",
]
MEDIA_STATUS_FIELDS = [
    "全株图片状态",
    "花图片状态",
    "果实图片状态",
    "叶图片状态",
    "茎图片状态",
]

FIELD_KEYS = {
    "花颜色": "color",
    "花序类型": "inflorescence",
    "花冠/花形": "corollaShape",
    "花对称性": "symmetry",
    "花性": "sexuality",
    "花瓣/花被片数": "tepalCount",
    "花大小": "size",
    "花期": "season",
    "花完整描述": "description",
    "果实类型": "type",
    "果实形状": "shape",
    "果实颜色": "color",
    "果实大小": "size",
    "果期": "season",
    "种子特征": "seed",
    "果实完整描述": "description",
    "茎质地": "texture",
    "茎姿态": "posture",
    "分枝方式": "branching",
    "茎形/结构": "structure",
    "茎颜色": "color",
    "茎毛被/刺": "surface",
    "植株高度": "height",
    "茎完整描述": "description",
    "叶序": "arrangement",
    "叶序细分": "arrangementDetail",
    "叶类型": "type",
    "复叶类型": "compoundType",
    "叶片形状": "shape",
    "叶缘": "margin",
    "叶尖": "apex",
    "叶基": "base",
    "叶脉": "venation",
    "叶质": "texture",
    "叶面/毛被": "surface",
    "叶柄": "petiole",
    "托叶/叶鞘": "stipuleSheath",
    "叶大小": "size",
    "叶完整描述": "description",
}


def clean_text(value: Any, *, allow_placeholder: bool = False) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    text = re.sub(r"\s+", " ", str(value).strip())
    if not allow_placeholder and text.lower() in PLACEHOLDERS:
        return ""
    return text


def safe_number(value: Any, default: float = 0) -> float:
    text = clean_text(value, allow_placeholder=True).replace("%", "")
    if not text:
        return default
    try:
        number = float(text)
        return number / 100 if "%" in str(value) else number
    except ValueError:
        return default


def read_sheet(name: str, header: int) -> pd.DataFrame:
    frame = pd.read_excel(WORKBOOK_PATH, sheet_name=name, header=header, dtype=object)
    frame = frame.dropna(how="all").reset_index(drop=True)
    frame.columns = [clean_text(column, allow_placeholder=True) for column in frame.columns]
    return frame


def keyed_values(row: pd.Series, fields: list[str]) -> dict[str, str]:
    return {FIELD_KEYS[field]: clean_text(row.get(field)) for field in fields}


def flatten_text(value: Any) -> list[str]:
    if isinstance(value, str):
        return [value] if value else []
    if isinstance(value, dict):
        return [item for child in value.values() for item in flatten_text(child)]
    if isinstance(value, list):
        return [item for child in value for item in flatten_text(child)]
    return []


def build_reference_media(serial: int) -> list[dict[str, str]]:
    prefix = f"P{serial:03d}-"
    files = sorted(
        path
        for path in REFERENCE_IMAGE_DIR.glob(f"{prefix}*")
        if path.is_file() and path.suffix.lower() in IMAGE_SUFFIXES
    )
    return [
        {
            "id": f"IPLANT-{serial:03d}-{index:02d}",
            "url": f"/plant-images/{path.name}",
            "thumbnailUrl": f"/plant-images/{path.name}",
            "sourceType": "iPlant参考图",
            "organLabel": "参考图",
        }
        for index, path in enumerate(files, start=1)
    ]


def taxonomy_review(plant_id: str, name: str, latin_name: str) -> tuple[str, str, list[dict[str, Any]], list[str]]:
    accepted_name = name
    accepted_latin = latin_name
    revisions: list[dict[str, Any]] = []
    flags: list[str] = []

    if plant_id == "HBFC-280":
        accepted_latin = "Polygonatum sibiricum"
        revisions.append(
            {
                "field": "latinName",
                "originalValue": latin_name,
                "acceptedValue": accepted_latin,
                "reviewStatus": "定向校订",
                "authorityUrl": "https://www.iplant.cn/bk/0AFDB1D075A2CECC",
                "note": "黄精与小黄花菜条目发生拉丁名串位，已按权威植物志定向修订。",
            }
        )
    elif plant_id == "HBFC-254":
        accepted_name = "芒𫈰草"
        revisions.append(
            {
                "field": "name",
                "originalValue": name,
                "acceptedValue": accepted_name,
                "reviewStatus": "乱码修复",
                "authorityUrl": "https://www.iplant.cn/info/Koeleria%20macrantha",
            }
        )
    elif plant_id == "HBFC-255":
        accepted_name = "𫈰草"
        accepted_latin = "Koeleria macrantha"
        revisions.extend(
            [
                {
                    "field": "name",
                    "originalValue": name,
                    "acceptedValue": accepted_name,
                    "reviewStatus": "乱码修复",
                    "authorityUrl": "https://www.iplant.cn/info/Koeleria%20macrantha",
                },
                {
                    "field": "latinName",
                    "originalValue": latin_name,
                    "acceptedValue": accepted_latin,
                    "reviewStatus": "接受名关系",
                    "authorityUrl": "https://www.iplant.cn/info/Koeleria%20macrantha",
                    "note": "记录 Koeleria cristata → Koeleria macrantha 的接受名关系。",
                },
            ]
        )

    if plant_id in {"HBFC-056", "HBFC-057"}:
        flags.append("Dontostemon integrifolius 疑似重复；保留原 ID，等待教师复核")
        revisions.append(
            {
                "field": "record",
                "originalValue": plant_id,
                "acceptedValue": plant_id,
                "reviewStatus": "疑似重复未合并",
                "authorityUrl": "https://www.iplant.cn/foc/pdf/Brassicaceae.pdf",
            }
        )

    return accepted_name, accepted_latin, revisions, flags


def split_suggestions(text: str) -> list[str]:
    return [item.strip() for item in re.split(r"[、，,；;/]+", text) if item.strip()]


def build_plants(feature_frame: pd.DataFrame, enhanced_frame: pd.DataFrame) -> list[dict[str, Any]]:
    plants: list[dict[str, Any]] = []
    enhanced_rows = enhanced_frame.to_dict("records")

    for index, row in feature_frame.iterrows():
        serial = index + 1
        enhanced = enhanced_rows[index] if index < len(enhanced_rows) else {}
        plant_id = clean_text(row.get("植物ID")) or f"HBFC-{serial:03d}"
        original_name = clean_text(row.get("中文名"))
        original_latin = clean_text(row.get("拉丁名"))
        accepted_name, accepted_latin, revisions, flags = taxonomy_review(
            plant_id, original_name, original_latin
        )

        alias = clean_text(enhanced.get("别名")) or clean_text(row.get("别名"))

        reference_media = build_reference_media(serial)
        completeness = safe_number(row.get("结构化完整度"))
        if completeness > 1:
            completeness /= 100

        plant: dict[str, Any] = {
            "id": plant_id,
            "legacyId": f"P{serial:03d}",
            "serial": serial,
            "names": {
                "chinese": accepted_name,
                "alias": alias,
                "latin": accepted_latin,
                "originalChinese": original_name if accepted_name != original_name else "",
                "originalLatin": original_latin if accepted_latin != original_latin else "",
            },
            "taxonomy": {
                "family": clean_text(row.get("科")) or clean_text(enhanced.get("科名")),
                "familyLatin": clean_text(enhanced.get("科拉丁名")),
                "genus": clean_text(row.get("属")) or clean_text(enhanced.get("属名")),
                "genusLatin": clean_text(enhanced.get("属拉丁名")),
            },
            "ecology": {
                "lifeForm": clean_text(row.get("生活型")) or clean_text(enhanced.get("生活型")),
                "habitat": clean_text(row.get("生境")) or clean_text(enhanced.get("生活习性/生境")),
            },
            "profile": {
                "introduction": clean_text(enhanced.get("植物外貌特征")) or clean_text(row.get("整体外貌")),
                "habitatNote": clean_text(enhanced.get("生活习性/生境")) or clean_text(row.get("生境")),
                "sourceIntroduction": clean_text(enhanced.get("iPlant补充介绍")),
            },
            "morphology": {
                "appearance": clean_text(row.get("整体外貌")) or clean_text(enhanced.get("植物外貌特征")),
                "flower": keyed_values(row, FLOWER_FIELDS),
                "fruit": keyed_values(row, FRUIT_FIELDS),
                "stem": keyed_values(row, STEM_FIELDS),
                "leaf": keyed_values(row, LEAF_FIELDS),
                "root": clean_text(row.get("根/地下器官")),
            },
            "identification": {
                "quickMethod": clean_text(row.get("快速识别方法")) or clean_text(enhanced.get("快速识别方法")),
                "keyCombination": clean_text(row.get("关键识别组合")),
                "suggestedParts": split_suggestions(clean_text(row.get("建议拍摄部位"))),
                "steps": [],
            },
            "media": {
                "localSamples": [],
                "iplantReferences": reference_media,
                "status": {
                    "whole": clean_text(row.get("全株图片状态")),
                    "flower": clean_text(row.get("花图片状态")),
                    "fruit": clean_text(row.get("果实图片状态")),
                    "leaf": clean_text(row.get("叶图片状态")),
                    "stem": clean_text(row.get("茎图片状态")),
                },
            },
            "quality": {
                "completeness": round(completeness, 3),
                "needsReview": clean_text(row.get("需人工复核")) not in {"", "否", "无需"},
                "reviewFlags": flags,
                "revisions": revisions,
            },
            "sources": {
                "workbook": WORKBOOK_PATH.name,
                "metadataSheet": "01植物部位特征",
                "iplantUrl": clean_text(enhanced.get("iPlant详情页")) or clean_text(row.get("iPlant来源")),
                "iplantQuery": clean_text(enhanced.get("iPlant查询键")),
            },
        }

        quick = plant["identification"]["quickMethod"]
        key_combo = plant["identification"]["keyCombination"]
        plant["identification"]["steps"] = [
            value
            for value in [
                f"先看整体：{plant['morphology']['appearance']}" if plant["morphology"]["appearance"] else "",
                f"再核对组合：{key_combo}" if key_combo else "",
                f"最后确认：{quick}" if quick else "",
            ]
            if value
        ]
        plant["searchText"] = " ".join(dict.fromkeys(flatten_text(plant))).lower()
        plants.append(plant)

    return plants


def build_feature_index(frame: pd.DataFrame) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for index, row in frame.iterrows():
        plant_ids = re.findall(r"HBFC-\d{3}", clean_text(row.get("植物ID清单")))
        names = [
            item.strip()
            for item in re.split(r"[、，,；;]+", clean_text(row.get("对应植物清单")))
            if item.strip()
        ]
        records.append(
            {
                "id": clean_text(row.get("特征ID")) or f"FEATURE-{index + 1:03d}",
                "organ": clean_text(row.get("部位")),
                "category": clean_text(row.get("特征类别")),
                "value": clean_text(row.get("特征值")),
                "beginnerExplanation": clean_text(row.get("新手解释")),
                "plantCount": int(safe_number(row.get("对应植物数量"))),
                "plantNames": names,
                "plantIds": plant_ids,
                "recommendedPart": clean_text(row.get("推荐拍摄部位")),
                "example": clean_text(row.get("图片示例")),
                "note": clean_text(row.get("备注")),
            }
        )
    return records


def build_glossary_and_checklist() -> tuple[list[dict[str, str]], list[dict[str, Any]]]:
    workbook = openpyxl.load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
    sheet = workbook["03术语与拍摄规范"]

    glossary = []
    for index, row in enumerate(sheet.iter_rows(min_row=5, max_row=40, values_only=True), start=1):
        values = [clean_text(value) for value in row[:6]]
        if not values[1]:
            continue
        glossary.append(
            {
                "id": f"TERM-{index:03d}",
                "organ": values[0],
                "term": values[1],
                "explanation": values[2],
                "captureMethod": values[3],
                "indexField": values[4],
                "namingExample": values[5],
            }
        )

    checklist = []
    for row in sheet.iter_rows(min_row=45, max_row=49, values_only=True):
        values = list(row[:6])
        checklist.append(
            {
                "id": f"CAPTURE-{int(values[0]):02d}",
                "subject": clean_text(values[1]),
                "minimumCount": int(safe_number(values[2])),
                "requirement": clean_text(values[3]),
                "commonIssue": clean_text(values[4]),
                "partCode": clean_text(values[5]),
            }
        )
    workbook.close()
    return glossary, checklist


def prepare_local_image(source: Path, target: Path, max_size: int, quality: int) -> None:
    with Image.open(source) as opened:
        image = ImageOps.exif_transpose(opened)
        if image.mode not in {"RGB", "RGBA"}:
            image = image.convert("RGB")
        image.thumbnail((max_size, max_size), Image.Resampling.LANCZOS)
        target.parent.mkdir(parents=True, exist_ok=True)
        image.save(target, "WEBP", quality=quality, method=6, exif=b"")


def import_local_samples(frame: pd.DataFrame, plants: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    (LOCAL_IMAGE_DIR / "display").mkdir(parents=True, exist_ok=True)
    (LOCAL_IMAGE_DIR / "thumb").mkdir(parents=True, exist_ok=True)

    plants_by_id = {plant["id"]: plant for plant in plants}
    samples: list[dict[str, Any]] = []
    pending: list[dict[str, Any]] = []

    for index, row in frame.iterrows():
        serial = index + 1
        sample_id = f"SAMPLE-{serial:04d}"
        relative_source = clean_text(row.get("相对路径"), allow_placeholder=True)
        source_path = SOURCE_ROOT / Path(relative_source.replace("\\", "/"))
        if not source_path.exists():
            raise FileNotFoundError(f"群聊图片不存在: {source_path}")

        file_name = f"S{serial:04d}.webp"
        display_path = LOCAL_IMAGE_DIR / "display" / file_name
        thumb_path = LOCAL_IMAGE_DIR / "thumb" / file_name
        if not display_path.exists() or display_path.stat().st_mtime < source_path.stat().st_mtime:
            prepare_local_image(source_path, display_path, 1600, 84)
        if not thumb_path.exists() or thumb_path.stat().st_mtime < source_path.stat().st_mtime:
            prepare_local_image(source_path, thumb_path, 560, 78)

        plant_id = clean_text(row.get("植物ID"))
        matched = clean_text(row.get("匹配状态")) == "已匹配01" and plant_id in plants_by_id
        sample = {
            "id": sample_id,
            "plantId": plant_id if matched else "",
            "submittedName": clean_text(row.get("植物种类")),
            "organLabel": clean_text(row.get("部位")) or "未标注",
            "matched": matched,
            "reviewStatus": "已挂接" if matched else "待复核",
            "displayUrl": f"/local-samples/display/{file_name}",
            "thumbnailUrl": f"/local-samples/thumb/{file_name}",
            "submittedAt": clean_text(row.get("群消息时间")),
            "timeMeaning": "提交时间",
            "sourceType": "实习群样本",
        }
        samples.append(sample)
        if matched:
            plants_by_id[plant_id]["media"]["localSamples"].append(sample)
        else:
            pending.append(sample)

    return samples, pending


def build_summary(
    plants: list[dict[str, Any]],
    feature_headers: list[str],
    feature_index: list[dict[str, Any]],
    glossary: list[dict[str, str]],
    checklist: list[dict[str, Any]],
    samples: list[dict[str, Any]],
    pending: list[dict[str, Any]],
) -> dict[str, Any]:
    reference_count = sum(len(plant["media"]["iplantReferences"]) for plant in plants)
    matched_count = sum(1 for sample in samples if sample["matched"])
    local_covered = {sample["plantId"] for sample in samples if sample["matched"]}
    pending_names = {sample["submittedName"] for sample in pending}
    aliases = [plant["names"]["alias"] for plant in plants if plant["names"]["alias"]]

    def count(field: str) -> list[dict[str, Any]]:
        counts = Counter(plant[field.split(".")[0]][field.split(".")[1]] or "待补充" for plant in plants)
        return [{"name": name, "count": total} for name, total in counts.most_common()]

    return {
        "recordCount": len(plants),
        "featureFieldCount": len(feature_headers),
        "reverseIndexCount": len(feature_index),
        "glossaryCount": len(glossary),
        "captureChecklistCount": len(checklist),
        "referenceImageCount": reference_count,
        "localImageCount": len(samples),
        "matchedLocalImages": matched_count,
        "pendingLocalImages": len(pending),
        "localCoveredPlants": len(local_covered),
        "pendingNameCount": len(pending_names),
        "aliasCount": len(aliases),
        "familyCount": len({plant["taxonomy"]["family"] for plant in plants if plant["taxonomy"]["family"]}),
        "genusCount": len({plant["taxonomy"]["genus"] for plant in plants if plant["taxonomy"]["genus"]}),
        "averageCompleteness": round(sum(plant["quality"]["completeness"] for plant in plants) / len(plants), 3),
        "lifeForms": count("ecology.lifeForm"),
        "families": count("taxonomy.family"),
        "dataVersion": "V2-2026-07-10",
        "privacy": {
            "locationMetadataRemoved": True,
            "chatIdentityRemoved": True,
            "localImageExifRemoved": True,
        },
    }


def write_json(name: str, value: Any) -> None:
    (PUBLIC_DATA_DIR / name).write_text(
        json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def validate_counts(summary: dict[str, Any]) -> None:
    expected = {
        "recordCount": 286,
        "featureFieldCount": 62,
        "reverseIndexCount": 284,
        "glossaryCount": 36,
        "captureChecklistCount": 5,
        "referenceImageCount": 785,
        "localImageCount": 215,
        "matchedLocalImages": 176,
        "pendingLocalImages": 39,
        "localCoveredPlants": 53,
        "pendingNameCount": 14,
        "aliasCount": 17,
    }
    mismatches = {
        key: {"expected": expected_value, "actual": summary.get(key)}
        for key, expected_value in expected.items()
        if summary.get(key) != expected_value
    }
    if mismatches:
        raise ValueError(f"数据校验失败: {json.dumps(mismatches, ensure_ascii=False)}")


def main() -> None:
    if not WORKBOOK_PATH.exists():
        raise FileNotFoundError(WORKBOOK_PATH)

    PUBLIC_DATA_DIR.mkdir(parents=True, exist_ok=True)
    feature_frame = read_sheet("01植物部位特征", header=3)
    enhanced_frame = read_sheet("iPlant增强名录", header=0)
    reverse_frame = read_sheet("02特征反向索引", header=3)
    group_frame = read_sheet("群聊图片索引", header=0)

    plants = build_plants(feature_frame, enhanced_frame)
    feature_index = build_feature_index(reverse_frame)
    glossary, checklist = build_glossary_and_checklist()
    samples, pending = import_local_samples(group_frame, plants)
    summary = build_summary(
        plants,
        list(feature_frame.columns),
        feature_index,
        glossary,
        checklist,
        samples,
        pending,
    )
    validate_counts(summary)

    write_json("plants.json", plants)
    write_json("featureIndex.json", feature_index)
    write_json("glossary.json", glossary)
    write_json("captureChecklist.json", checklist)
    write_json("pendingSamples.json", pending)
    write_json("catalogSummary.json", summary)

    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
